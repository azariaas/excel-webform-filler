"""Create a local demo website and a fictitious leads spreadsheet."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
DEMO_DIR = ROOT / "demo_site"
DATA_DIR = ROOT / "data"

STORES = {
    "NORTH": "Northside Store",
    "SOUTH": "Southside Store",
    "EAST": "Eastside Store",
}

PRODUCTS = {
    "desk-lamp": "Desk Lamp",
    "office-chair": "Office Chair",
    "standing-desk": "Standing Desk",
    "bookshelf": "Bookshelf",
    "monitor-arm": "Monitor Arm",
}

FORM_HTML = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>{product} quote — {store}</title>
  <style>
    body {{ font-family: Arial, sans-serif; max-width: 520px; margin: 40px auto; }}
    label, input, select, textarea, button {{ display: block; width: 100%; margin-bottom: 12px; }}
    textarea {{ min-height: 80px; }}
    .pref {{ display: flex; gap: 8px; align-items: center; }}
    .pref input {{ width: auto; }}
  </style>
</head>
<body>
  <h1>{store}</h1>
  <h2>Quote request: {product}</h2>
  <form action="/thanks.html" method="get">
    <label>Version
      <select name="version">
        <option>Standard</option>
        <option>Plus</option>
        <option>Premium</option>
      </select>
    </label>
    <input name="name" placeholder="Full name" required>
    <input name="phone" placeholder="Phone">
    <input name="email" placeholder="Email" type="email">
    <textarea name="notes" placeholder="Questions or notes"></textarea>
    <div class="pref">
      <input type="radio" name="contact_pref" value="email" id="pref-email">
      <label for="pref-email">Email</label>
    </div>
    <div class="pref">
      <input type="radio" name="contact_pref" value="whatsapp" id="pref-whatsapp">
      <label for="pref-whatsapp">WhatsApp</label>
    </div>
    <div class="pref">
      <input type="checkbox" name="privacy" id="privacy">
      <label for="privacy">I agree to the privacy policy</label>
    </div>
    <button type="submit">Send quote request</button>
  </form>
</body>
</html>
"""

THANKS_HTML = """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <title>Request received</title>
</head>
<body style="font-family: Arial, sans-serif; max-width: 520px; margin: 40px auto;">
  <h1>Request received</h1>
  <p>This is a local demo page. No data was sent to an external site.</p>
</body>
</html>
"""


def write_demo_site() -> None:
    for code, store in STORES.items():
        for slug, product in PRODUCTS.items():
            folder = DEMO_DIR / "stores" / code.lower() / "new"
            folder.mkdir(parents=True, exist_ok=True)
            (folder / f"{slug}.html").write_text(
                FORM_HTML.format(store=store, product=product),
                encoding="utf-8",
            )
    DEMO_DIR.mkdir(parents=True, exist_ok=True)
    (DEMO_DIR / "thanks.html").write_text(THANKS_HTML, encoding="utf-8")
    (DEMO_DIR / "index.html").write_text(
        "<h1>Harbor Home demo</h1><p>Start the form filler while this site is being served.</p>",
        encoding="utf-8",
    )


def write_spreadsheet() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    samples = {
        "NORTH": [
            ("Casey Quinn", "555-0101", "casey.quinn@example.com", "Desk Lamp", "Plus", "Need a spare bulb"),
            ("Sam Ortega", "555-0102", "sam.ortega@example.com", "Office Chair", "Standard", ""),
        ],
        "SOUTH": [
            ("Riley Chen", "555-0103", "riley.chen@example.com", "Standing Desk", "Premium", "Prefer oak finish"),
            ("Morgan Ellis", "555-0104", "morgan.ellis@example.com", "Bookshelf", "Standard", ""),
        ],
        "EAST": [
            ("Taylor Brooks", "555-0105", "taylor.brooks@example.com", "Monitor Arm", "Plus", ""),
            ("Jamie Patel", "555-0106", "jamie.patel@example.com", "Unknown Item", "Standard", "Should be skipped"),
        ],
    }

    first = True
    header_font = Font(bold=True)
    yellow = PatternFill("solid", fgColor="FFFF00")

    for sheet_name, rows in samples.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = sheet_name
        ws["A1"] = "Harbor Home — sample quote requests"
        ws["A2"] = "Fictitious data for local demonstration only"
        headers = ["NAME", "PHONE", "EMAIL", "PRODUCT", "VERSION", "NOTES"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(3, col, header)
            cell.font = header_font
        for offset, row in enumerate(rows):
            for col, value in enumerate(row, start=1):
                ws.cell(4 + offset, col, value)

    # Pretend the first Northside row was already submitted
    for cell in wb["NORTH"][4]:
        cell.fill = yellow

    example = DATA_DIR / "leads.example.xlsx"
    working = DATA_DIR / "leads.xlsx"
    wb.save(example)
    wb.save(working)


def main() -> None:
    write_demo_site()
    write_spreadsheet()
    print(f"Demo site: {DEMO_DIR}")
    print(f"Spreadsheet: {DATA_DIR / 'leads.xlsx'}")


if __name__ == "__main__":
    main()
