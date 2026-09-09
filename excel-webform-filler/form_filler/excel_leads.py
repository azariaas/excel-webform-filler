from __future__ import annotations

from dataclasses import dataclass

from openpyxl.cell.cell import Cell
from openpyxl.styles import PatternFill
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

@dataclass(frozen=True)
class LeadRow:
    sheet: str
    row_number: int
    name: str
    phone: str
    email: str
    product: str
    version: str
    notes: str


def fill_color_rgb(cell: Cell) -> str:
    fill = cell.fill
    if not fill or not fill.fgColor:
        return ""
    if fill.fgColor.type != "rgb":
        return ""
    rgb = str(fill.fgColor.rgb or "").upper().lstrip("#")
    if len(rgb) >= 6:
        return rgb[-6:]
    return rgb


def is_processed_fill(cell: Cell, processed_fill: str = "FFFF00") -> bool:
    rgb = fill_color_rgb(cell)
    if not rgb:
        return False
    target = processed_fill.upper().lstrip("#")
    if len(target) >= 6:
        target = target[-6:]
    return rgb == target


def header_map(ws: Worksheet, header_row: int) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for cell in ws[header_row]:
        if cell.value:
            mapping[str(cell.value).strip().upper()] = cell.column
    return mapping


def _cell_text(ws: Worksheet, row: int, column: int | None) -> str:
    if not column:
        return ""
    value = ws.cell(row=row, column=column).value
    return str(value).strip() if value is not None else ""


def read_lead_rows(
    wb: Workbook,
    locations: dict[str, str],
    header_row: int,
    data_start_row: int,
    processed_fill: str,
) -> tuple[list[LeadRow], list[str]]:
    leads: list[LeadRow] = []
    skipped_sheets: list[str] = []

    for sheet_name in wb.sheetnames:
        location_key = sheet_name.strip().upper()
        if location_key not in locations:
            skipped_sheets.append(sheet_name)
            continue

        ws = wb[sheet_name]
        columns = header_map(ws, header_row)
        name_col = columns.get("NAME")
        if not name_col:
            raise KeyError(f"Sheet '{sheet_name}' is missing a NAME column on row {header_row}")

        for row_cells in ws.iter_rows(min_row=data_start_row):
            row_number = row_cells[0].row
            name_cell = ws.cell(row=row_number, column=name_col)
            name = _cell_text(ws, row_number, name_col)
            if not name:
                continue
            if is_processed_fill(name_cell, processed_fill):
                continue

            leads.append(
                LeadRow(
                    sheet=sheet_name,
                    row_number=row_number,
                    name=name,
                    phone=_cell_text(ws, row_number, columns.get("PHONE")),
                    email=_cell_text(ws, row_number, columns.get("EMAIL")),
                    product=_cell_text(ws, row_number, columns.get("PRODUCT")).upper(),
                    version=_cell_text(ws, row_number, columns.get("VERSION")),
                    notes=_cell_text(ws, row_number, columns.get("NOTES")),
                )
            )

    return leads, skipped_sheets


def mark_row_processed(ws: Worksheet, row_number: int, fill_hex: str) -> None:
    fill = PatternFill("solid", fgColor=fill_hex.lstrip("#"))
    for cell in ws[row_number]:
        cell.fill = fill
