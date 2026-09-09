from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from form_filler.excel_leads import is_processed_fill, read_lead_rows
from form_filler.urls import build_form_url, lookup_product_slug


def test_lookup_product_slug_normalizes_spaces_and_case():
    catalog = {"DESK LAMP": "desk-lamp"}
    assert lookup_product_slug(" desk   lamp ", catalog) == "desk-lamp"
    assert lookup_product_slug("Unknown", catalog) is None


def test_build_form_url_joins_location_and_slug():
    url = build_form_url("http://127.0.0.1:8000", "/stores/north/new/", "desk-lamp")
    assert url == "http://127.0.0.1:8000/stores/north/new/desk-lamp.html"


def test_is_processed_fill_detects_yellow():
    wb = Workbook()
    cell = wb.active["A1"]
    cell.fill = PatternFill("solid", fgColor="FFFF00")
    assert is_processed_fill(cell) is True
    wb.active["A2"].value = "plain"
    assert is_processed_fill(wb.active["A2"]) is False


def _workbook_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "NORTH"
    ws["A3"] = "NAME"
    ws["B3"] = "PHONE"
    ws["C3"] = "EMAIL"
    ws["D3"] = "PRODUCT"
    ws["E3"] = "VERSION"
    ws["F3"] = "NOTES"
    ws["A4"] = "Casey Quinn"
    ws["D4"] = "Desk Lamp"
    ws["A5"] = "Already Sent"
    ws["A5"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["A6"] = ""
    extra = wb.create_sheet("UNKNOWN")
    extra["A3"] = "NAME"
    extra["A4"] = "Should Skip Sheet"
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_read_lead_rows_skips_yellow_blank_and_unknown_sheets():
    wb = load_workbook(BytesIO(_workbook_bytes()))
    leads, skipped = read_lead_rows(
        wb,
        locations={"NORTH": "/stores/north/new/"},
        header_row=3,
        data_start_row=4,
        processed_fill="FFFF00",
    )
    assert skipped == ["UNKNOWN"]
    assert [lead.name for lead in leads] == ["Casey Quinn"]
    assert leads[0].product == "DESK LAMP"
